"""Bước 7 - Sinh từ điển dữ liệu (data dictionary) cho các bảng đầu vào.

Phần mô tả nghiệp vụ được viết tay trong chính file này; phần thống kê (kiểu dữ
liệu, tỷ lệ khuyết, số giá trị phân biệt, miền giá trị, ví dụ) được đọc trực
tiếp từ các file parquet trong ``data/raw`` để bảo đảm từ điển luôn khớp với dữ
liệu thực tế. Danh mục mã lấy từ ``creditreview.data.reference``.

Kết quả:
    Tu_dien_du_lieu_bang_dau_vao.xlsx   8 sheet: tổng quan, 6 bảng, danh mục mã

Cách dùng:
    py scripts/07_export_data_dictionary.py
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT_DIR = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT_DIR / "data" / "raw"
OUT = ROOT_DIR / "Tu_dien_du_lieu_bang_dau_vao.xlsx"

FONT = "Arial"
C_HDR = "1F3864"      # xanh đậm
C_HDR_TXT = "FFFFFF"
C_TITLE = "1F3864"
C_BAND = "EDF2F9"
C_INTERNAL = "FFF2CC"  # vàng nhạt cho cột nội bộ
C_KEY = "E2EFDA"       # xanh nhạt cho khóa

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ---------------------------------------------------------------------------
# Mô tả các bảng
# ---------------------------------------------------------------------------
TABLES = [
    ("customers", "Khách hàng",
     "customer_id", "—", "1 dòng / khách hàng",
     "Hồ sơ nhân khẩu học và thông tin tài chính khai báo của khách hàng vay."),
    ("accounts", "Tài khoản",
     "account_id", "customer_id → customers", "1 dòng / tài khoản",
     "Tài khoản thanh toán, tiết kiệm, trả lương và thẻ của khách hàng tại tổ chức tín dụng."),
    ("transactions", "Giao dịch",
     "txn_id", "customer_id → customers; account_id → accounts", "1 dòng / giao dịch",
     "Lịch sử giao dịch tài khoản — nguồn dữ liệu của toàn bộ nhóm đặc trưng hành vi."),
    ("applications", "Hồ sơ vay",
     "application_id", "customer_id → customers", "1 dòng / hồ sơ vay",
     "Đề nghị cấp tín dụng. Trường apply_time là mốc cắt thời gian cho mọi đặc trưng."),
    ("cic_reports", "Bản tin CIC",
     "report_id", "application_id → applications; customer_id → customers", "1 dòng / lần truy vấn CIC",
     "Thông tin lịch sử tín dụng truy vấn từ Trung tâm Thông tin tín dụng Quốc gia."),
    ("labels", "Nhãn rủi ro",
     "application_id", "application_id → applications; customer_id → customers", "1 dòng / hồ sơ vay",
     "Kết quả trả nợ sau cửa sổ quan sát 12 tháng — biến mục tiêu của bài toán."),
]

INTERNAL_NOTE = ("CỘT NỘI BỘ của chương trình sinh dữ liệu. KHÔNG được dùng làm "
                 "đầu vào mô hình (gây rò rỉ nhãn); chỉ dùng để kiểm chứng và tính trần Bayes.")

# (mô tả, đơn vị, khóa/ràng buộc, bắt buộc)
DESC: dict[str, dict[str, tuple[str, str, str, str]]] = {
    "customers": {
        "customer_id": ("Mã định danh khách hàng trong hệ thống, dạng CUSnnnnnnnn", "", "Khóa chính", "Có"),
        "cif_number": ("Số CIF — mã khách hàng duy nhất trên hệ thống core banking", "", "Duy nhất", "Có"),
        "age": ("Tuổi của khách hàng tại thời điểm trích xuất dữ liệu", "năm", "", "Có"),
        "gender": ("Giới tính theo giấy tờ tùy thân", "", "", "Có"),
        "marital_status": ("Tình trạng hôn nhân khai báo", "", "", "Có"),
        "num_dependents": ("Số người phụ thuộc — dùng để ước tính chi phí sinh hoạt tối thiểu", "người", "", "Có"),
        "education_code": ("Mã trình độ học vấn (xem sheet Danh mục mã)", "", "1–5", "Có"),
        "occupation": ("Nhóm nghề nghiệp hiện tại", "", "", "Có"),
        "employer_type": ("Loại hình tổ chức nơi khách hàng làm việc", "", "", "Có"),
        "job_tenure_months": ("Thâm niên tại công việc hiện tại", "tháng", "", "Có"),
        "monthly_income": ("Thu nhập bình quân tháng khai báo — mẫu số của DTI/PTI", "VND", "", "Có"),
        "province": ("Tỉnh/thành phố nơi cư trú", "", "30 tỉnh/thành", "Có"),
        "residence_type": ("Hình thức sở hữu nơi ở hiện tại", "", "", "Có"),
        "customer_since_months": ("Số tháng kể từ khi thiết lập quan hệ với tổ chức tín dụng", "tháng", "", "Có"),
        "has_salary_account": ("Có tài khoản nhận lương tại tổ chức tín dụng hay không", "0/1", "", "Có"),
        "created_at": ("Thời điểm tạo bản ghi khách hàng", "", "", "Có"),
        "updated_at": ("Thời điểm cập nhật bản ghi gần nhất", "", "", "Có"),
        "_z": ("Biến tiềm ẩn biểu diễn mức tín nhiệm nội tại, phân phối N(0,1). " + INTERNAL_NOTE, "", "Nội bộ", "—"),
        "_income_volatility": ("Hệ số biến động thu nhập theo nhóm nghề. " + INTERNAL_NOTE, "", "Nội bộ", "—"),
        "_income_stability_base": ("Mức ổn định thu nhập cơ sở theo loại hình tổ chức. " + INTERNAL_NOTE, "", "Nội bộ", "—"),
        "_is_thin_file": ("Cờ khách hàng thiếu lịch sử tín dụng. " + INTERNAL_NOTE, "0/1", "Nội bộ", "—"),
    },
    "accounts": {
        "account_id": ("Mã tài khoản, dạng ACCnnnnnnnn", "", "Khóa chính", "Có"),
        "customer_id": ("Khách hàng sở hữu tài khoản", "", "Khóa ngoại → customers", "Có"),
        "account_type": ("Loại tài khoản: thanh toán, tiết kiệm, trả lương, thẻ", "", "", "Có"),
        "currency": ("Loại tiền của tài khoản", "", "", "Có"),
        "open_date": ("Ngày mở tài khoản", "", "", "Có"),
        "avg_balance_3m": ("Số dư bình quân 3 tháng gần nhất", "VND", "", "Có"),
        "current_balance": ("Số dư tại thời điểm trích xuất", "VND", "", "Có"),
        "status": ("Trạng thái hoạt động của tài khoản", "", "", "Có"),
    },
    "transactions": {
        "txn_id": ("Số thứ tự giao dịch", "", "Khóa chính", "Có"),
        "customer_id": ("Khách hàng phát sinh giao dịch — khóa phân vùng khi tính đặc trưng", "", "Khóa ngoại → customers", "Có"),
        "account_id": ("Tài khoản phát sinh giao dịch", "", "Khóa ngoại → accounts", "Có"),
        "txn_time": ("Thời điểm phát sinh giao dịch. Chỉ giao dịch có txn_time < apply_time "
                     "mới được dùng để tính đặc trưng (đúng đắn theo thời điểm)", "", "", "Có"),
        "amount": ("Giá trị giao dịch, luôn là số dương; chiều tiền nằm ở trường direction", "VND", "> 0", "Có"),
        "direction": ("Chiều dòng tiền: CR = tiền vào, DR = tiền ra", "", "", "Có"),
        "category": ("Nhóm mục đích giao dịch (xem sheet Danh mục mã)", "", "18 nhóm", "Có"),
        "channel": ("Kênh thực hiện giao dịch", "", "7 kênh", "Có"),
    },
    "applications": {
        "application_id": ("Mã hồ sơ vay, dạng APPnnnnnnnn", "", "Khóa chính", "Có"),
        "customer_id": ("Khách hàng đứng tên hồ sơ", "", "Khóa ngoại → customers", "Có"),
        "apply_time": ("Thời điểm nộp hồ sơ. Đây là MỐC CẮT THỜI GIAN: mọi đặc trưng chỉ được "
                       "tính từ dữ liệu phát sinh trước thời điểm này", "", "", "Có"),
        "product_code": ("Mã sản phẩm tín dụng (xem sheet Danh mục mã)", "", "6 sản phẩm", "Có"),
        "requested_amount": ("Số tiền đề nghị vay", "VND", "", "Có"),
        "tenor_months": ("Kỳ hạn vay đề nghị", "tháng", "", "Có"),
        "interest_rate": ("Lãi suất danh nghĩa năm, lưu dưới dạng thập phân (0,182 = 18,2%/năm)", "tỷ lệ/năm", "", "Có"),
        "loan_purpose": ("Mục đích sử dụng vốn khai báo", "", "9 nhóm", "Có"),
        "channel": ("Kênh tiếp nhận hồ sơ", "", "5 kênh", "Có"),
        "has_collateral": ("Hồ sơ có tài sản bảo đảm hay không", "0/1", "", "Có"),
        "collateral_value": ("Giá trị định giá tài sản bảo đảm; bằng 0 khi vay tín chấp", "VND", "≥ 0", "Có"),
        "ltv": ("Tỷ lệ cho vay trên giá trị tài sản bảo đảm (loan-to-value). ĐỂ TRỐNG với hồ sơ "
                "vay tín chấp (has_collateral = 0) — giá trị khuyết ở đây mang ý nghĩa nghiệp vụ, "
                "không phải lỗi dữ liệu", "tỷ lệ", "≈0,28–1,09", "Không (chỉ có với hồ sơ có TSBĐ)"),
        "monthly_installment": ("Số tiền phải trả hàng tháng, tính theo công thức niên kim từ "
                                "số tiền vay, lãi suất và kỳ hạn — tử số của PTI", "VND", "", "Có"),
        "_prod_risk_offset": ("Hệ số rủi ro cơ sở của sản phẩm. " + INTERNAL_NOTE, "", "Nội bộ", "—"),
        "_prod_secured": ("Cờ sản phẩm có tài sản bảo đảm. " + INTERNAL_NOTE, "0/1", "Nội bộ", "—"),
    },
    "cic_reports": {
        "report_id": ("Mã bản tin CIC", "", "Khóa chính", "Có"),
        "application_id": ("Hồ sơ vay đã kích hoạt lần truy vấn này", "", "Khóa ngoại → applications", "Có"),
        "customer_id": ("Khách hàng được truy vấn", "", "Khóa ngoại → customers", "Có"),
        "query_time": ("Thời điểm truy vấn CIC", "", "", "Có"),
        "valid_until": ("Thời hạn hiệu lực nghiệp vụ của bản tin — cơ sở xác định thời hạn lưu đệm (cache TTL)", "", "", "Có"),
        "cic_score": ("Điểm tín dụng do CIC cung cấp", "điểm", "300–850", "Không (thin-file)"),
        "current_group": ("Nhóm nợ hiện tại theo Thông tư 11/2021/TT-NHNN", "", "1–5", "Không (thin-file)"),
        "worst_group_12m": ("Nhóm nợ xấu nhất trong 12 tháng gần nhất", "", "1–5", "Không (thin-file)"),
        "dpd_max_12m": ("Số ngày quá hạn lớn nhất trong 12 tháng gần nhất; ≥ 90 ngày là nợ xấu", "ngày", "≥ 0", "Không (thin-file)"),
        "num_institutions": ("Số tổ chức tín dụng khách hàng đang có quan hệ", "tổ chức", "≥ 0", "Có"),
        "num_active_contracts": ("Số hợp đồng tín dụng còn hiệu lực", "hợp đồng", "≥ 0", "Có"),
        "num_closed_contracts": ("Số hợp đồng tín dụng đã tất toán", "hợp đồng", "≥ 0", "Có"),
        "total_outstanding": ("Tổng dư nợ hiện tại tại toàn bộ tổ chức tín dụng", "VND", "≥ 0", "Không (thin-file)"),
        "total_credit_limit": ("Tổng hạn mức tín dụng được cấp", "VND", "≥ 0", "Không (thin-file)"),
        "credit_utilization": ("Tỷ lệ sử dụng hạn mức = tổng dư nợ / tổng hạn mức", "tỷ lệ", "0–1", "Không (thin-file + khuyết ngẫu nhiên 1,4%)"),
        "monthly_debt_payment": ("Nghĩa vụ trả nợ hàng tháng tại các tổ chức khác — tử số của DTI", "VND", "≥ 0", "Không (thin-file)"),
        "num_inquiries_3m": ("Số lần hồ sơ bị truy vấn tín dụng trong 3 tháng — dấu hiệu khách hàng đang tìm vốn nhiều nơi", "lần", "0–15", "Có"),
        "has_written_off_debt": ("Có khoản nợ đã được xử lý rủi ro hay không", "0/1", "", "Có"),
        "oldest_account_months": ("Tuổi của quan hệ tín dụng lâu đời nhất", "tháng", "≥ 0", "Không (thin-file + khuyết ngẫu nhiên 2,2%)"),
        "num_bad_debt_institutions": ("Số tổ chức tín dụng mà khách hàng đang có nợ xấu", "tổ chức", "≥ 0", "Không (thin-file)"),
        "is_thin_file": ("Cờ khách hàng thiếu lịch sử tín dụng. Khi bằng 1, toàn bộ chỉ tiêu CIC "
                         "bị khuyết — đây là tình huống mô hình phải xử lý được", "0/1", "", "Có"),
        "_q": ("Phép đo có nhiễu của biến tiềm ẩn _z, dùng để sinh các chỉ tiêu CIC. " + INTERNAL_NOTE, "", "Nội bộ", "—"),
    },
    "labels": {
        "application_id": ("Hồ sơ vay được gán nhãn", "", "Khóa chính; khóa ngoại → applications", "Có"),
        "customer_id": ("Khách hàng của hồ sơ", "", "Khóa ngoại → customers", "Có"),
        "observation_time": ("Mốc bắt đầu cửa sổ quan sát, bằng apply_time của hồ sơ", "", "", "Có"),
        "performance_end": ("Mốc kết thúc cửa sổ quan sát = observation_time + 12 tháng", "", "", "Có"),
        "months_to_default": ("Số tháng kể từ khi cấp tín dụng tới thời điểm phát sinh nợ xấu; "
                              "để trống nếu hồ sơ không vỡ nợ", "tháng", "", "Không"),
        "default_flag": ("BIẾN MỤC TIÊU. Bằng 1 khi khoản vay rơi vào nhóm nợ 3–5 "
                         "(quá hạn từ 90 ngày) trong cửa sổ quan sát 12 tháng", "0/1", "", "Có"),
        "_pd_true": ("Xác suất vỡ nợ thật do chương trình sinh dữ liệu ấn định — cơ sở tính trần Bayes. "
                     + INTERNAL_NOTE, "xác suất", "Nội bộ", "—"),
        "_eta": ("Log-odds thật tương ứng với _pd_true. " + INTERNAL_NOTE, "log-odds", "Nội bộ", "—"),
        "risk_status": ("Trạng thái rủi ro: GOOD = không phát sinh nợ xấu; BAD = phát sinh nợ xấu; "
                        "INDETERMINATE = chưa đi hết cửa sổ quan sát nên chưa kết luận được",
                        "", "GOOD | BAD | INDETERMINATE", "Có"),
        "is_modelable": ("Hồ sơ có được đưa vào huấn luyện/đánh giá hay không. Bằng False khi "
                         "risk_status = INDETERMINATE", "", "True | False", "Có"),
    },
}

# ---------------------------------------------------------------------------
# Danh mục mã
# ---------------------------------------------------------------------------
import sys
sys.path.insert(0, str(ROOT_DIR / "src"))
from creditreview.data import reference as R  # noqa: E402

def w(x: float, nd: int = 3) -> str:
    """Trọng số/tỷ lệ theo quy ước dấu phẩy thập phân tiếng Việt."""
    return f"{x:.{nd}f}".replace(".", ",")


CODE_LISTS: list[tuple[str, str, list[tuple[str, str]]]] = [
    ("education_code", "customers.education_code",
     [(str(c), name) for c, name, *_ in R.EDUCATION]),
    ("occupation", "customers.occupation",
     [(name, f"trọng số {w(wt)}") for name, wt, *_ in R.OCCUPATIONS]),
    ("employer_type", "customers.employer_type",
     [(name, f"trọng số {w(wt)}") for name, wt, *_ in R.EMPLOYER_TYPES]),
    ("marital_status", "customers.marital_status",
     [(name, f"trọng số {w(wt)}") for name, wt, *_ in R.MARITAL_STATUS]),
    ("residence_type", "customers.residence_type",
     [(name, f"trọng số {w(wt)}") for name, wt, *_ in R.RESIDENCE_TYPES]),
    ("province", "customers.province",
     [(name, f"trọng số {w(wt)}") for name, wt, *_ in R.PROVINCES]),
    ("product_code", "applications.product_code",
     [(p["code"], f"{p['name']} — lãi suất {w(p['rate']*100, 1)}%/năm, "
       f"kỳ hạn {'/'.join(str(t) for t in p['tenors'])} tháng, "
       f"{'có' if p['secured'] else 'không có'} TSBĐ")
      for p in R.PRODUCTS]),
    ("loan_purpose", "applications.loan_purpose",
     [(name, f"trọng số {w(wt)}") for name, wt in R.LOAN_PURPOSES]),
    ("channel (hồ sơ)", "applications.channel",
     [(name, f"trọng số {w(wt)}") for name, wt, _ in R.CHANNELS]),
    ("category", "transactions.category",
     [(name, f"{'tiền vào' if d == 'CR' else 'tiền ra'} — trọng số {w(wt)}")
      for name, wt, d, _ in R.TXN_CATEGORIES]),
    ("channel (giao dịch)", "transactions.channel",
     [(name, f"trọng số {w(wt)}") for name, wt in R.TXN_CHANNELS]),
    ("nhóm nợ", "cic_reports.current_group / worst_group_12m",
     [(str(k), v) for k, v in R.DEBT_GROUPS.items()]),
    ("risk_status", "labels.risk_status",
     [("GOOD", "Không phát sinh nợ xấu trong cửa sổ quan sát"),
      ("BAD", "Phát sinh nợ xấu (nhóm 3–5 / quá hạn ≥ 90 ngày)"),
      ("INDETERMINATE", "Chưa đủ cửa sổ quan sát — loại khỏi tập mô hình")]),
]


# ---------------------------------------------------------------------------
# Tiện ích
# ---------------------------------------------------------------------------
def fmt_value(v) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    if isinstance(v, (pd.Timestamp, np.datetime64)):
        return pd.Timestamp(v).strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, (int, np.integer)):
        return f"{int(v):,}".replace(",", ".")
    if isinstance(v, (float, np.floating)):
        f = float(v)
        if abs(f) >= 1000:
            return f"{f:,.0f}".replace(",", ".")
        return f"{f:,.4f}".rstrip("0").rstrip(".").replace(".", ",")
    return str(v)


def dtype_vn(s: pd.Series) -> str:
    d = str(s.dtype)
    if d.startswith("datetime"):
        return "Ngày giờ"
    if d.startswith("int"):
        return "Số nguyên"
    if d.startswith("float"):
        return "Số thực"
    if d == "bool":
        return "Luận lý"
    return "Chuỗi ký tự"


KEY_MARKERS = ("Khóa", "Duy nhất", "Nội bộ")


def split_key_domain(raw: str) -> tuple[str, str]:
    """Tách phần khai báo thứ ba thành (khóa/ràng buộc, miền giá trị).

    Cùng một ô trong bảng khai báo có thể chứa thông tin khóa ("Khóa chính",
    "Khóa ngoại → customers") hoặc miền giá trị ("300–850", "1–5"); hai loại này
    thuộc hai cột khác nhau trong từ điển.
    """
    if not raw:
        return "", ""
    if any(raw.startswith(m) for m in KEY_MARKERS):
        return raw, ""
    return "", raw


def domain_of(s: pd.Series, declared: str) -> str:
    if declared:
        return declared
    nun = s.nunique(dropna=True)
    if s.dtype == object or str(s.dtype) == "bool":
        if nun <= 12:
            vals = sorted(str(x) for x in s.dropna().unique())
            return " | ".join(vals)
        return f"{nun:,} giá trị".replace(",", ".")
    return ""


# ---------------------------------------------------------------------------
# Dựng workbook
# ---------------------------------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = "Tổng quan"

HDRS = ["STT", "Trường dữ liệu", "Kiểu dữ liệu", "Khóa / Ràng buộc", "Bắt buộc",
        "Đơn vị", "Mô tả nghiệp vụ", "Miền giá trị / Danh mục", "Tỷ lệ khuyết",
        "Số giá trị phân biệt", "Giá trị nhỏ nhất", "Giá trị lớn nhất", "Ví dụ"]
WIDTHS = [5, 24, 12, 26, 20, 10, 62, 34, 11, 13, 22, 22, 22]


def style_header(sheet, row: int, ncol: int) -> None:
    for c in range(1, ncol + 1):
        cell = sheet.cell(row=row, column=c)
        cell.font = Font(name=FONT, size=10, bold=True, color=C_HDR_TXT)
        cell.fill = PatternFill("solid", fgColor=C_HDR)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def put_title(sheet, text: str, sub: str, ncol: int) -> None:
    sheet.cell(row=1, column=1, value=text).font = Font(name=FONT, size=13, bold=True, color=C_TITLE)
    sheet.cell(row=2, column=1, value=sub).font = Font(name=FONT, size=9, italic=True, color="595959")


meta = {}
for name, *_ in TABLES:
    df = pd.read_parquet(RAW_DIR / f"{name}.parquet")
    meta[name] = df

# ---- các sheet chi tiết ----------------------------------------------------
for name, vn, pk, fk, grain, purpose in TABLES:
    df = meta[name]
    sh = wb.create_sheet(name)
    put_title(sh, f"{name} — {vn}",
              f"{purpose}  |  Độ mịn: {grain}  |  Số dòng: {len(df):,}".replace(",", "."), len(HDRS))
    for i, h in enumerate(HDRS, start=1):
        sh.cell(row=3, column=i, value=h)
    style_header(sh, 3, len(HDRS))

    r = 4
    for stt, col in enumerate(df.columns, start=1):
        s = df[col]
        d, unit, raw_key, req = DESC[name].get(col, ("", "", "", ""))
        key, declared_domain = split_key_domain(raw_key)
        internal = col.startswith("_")
        if internal and not key:
            key = "Nội bộ"
        # "0/1" và "True/False" là tập giá trị hợp lệ, không phải đơn vị đo.
        if unit in ("0/1", "True/False"):
            declared_domain = declared_domain or unit
            unit = ""
        nn = s.notna()
        example = fmt_value(s[nn].iloc[0]) if nn.any() else ""
        if pd.api.types.is_numeric_dtype(s) and not pd.api.types.is_bool_dtype(s):
            vmin, vmax = fmt_value(s.min()), fmt_value(s.max())
        elif pd.api.types.is_datetime64_any_dtype(s):
            vmin, vmax = fmt_value(s.min()), fmt_value(s.max())
        else:
            vmin = vmax = ""

        row = [stt, col, dtype_vn(s), key, req, unit, d,
               domain_of(s, declared_domain), float(s.isna().mean()), int(s.nunique(dropna=True)),
               vmin, vmax, example]
        for i, v in enumerate(row, start=1):
            cell = sh.cell(row=r, column=i, value=v)
            cell.font = Font(name=FONT, size=10, italic=internal)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=i in (7, 8, 4),
                horizontal="center" if i in (1, 3, 5, 9, 10) else "left",
            )
        sh.cell(row=r, column=9).number_format = "0,00%"
        sh.cell(row=r, column=10).number_format = "#.##0"
        fill = C_INTERNAL if internal else (C_KEY if "Khóa" in key else (C_BAND if stt % 2 == 0 else None))
        if fill:
            for i in range(1, len(HDRS) + 1):
                sh.cell(row=r, column=i).fill = PatternFill("solid", fgColor=fill)
        r += 1

    note = ("Ghi chú: dòng in nghiêng nền vàng là cột nội bộ của chương trình sinh dữ liệu — "
            "không dùng làm đầu vào mô hình. Thống kê đọc trực tiếp từ data/raw/%s.parquet." % name)
    sh.cell(row=r + 1, column=1, value=note).font = Font(name=FONT, size=9, italic=True, color="595959")

    for i, w in enumerate(WIDTHS, start=1):
        sh.column_dimensions[get_column_letter(i)].width = w
    sh.freeze_panes = "C4"
    sh.auto_filter.ref = f"A3:{get_column_letter(len(HDRS))}{r - 1}"

# ---- sheet Tổng quan -------------------------------------------------------
OH = ["STT", "Bảng", "Tên tiếng Việt", "Số dòng", "Số trường", "Số trường có dữ liệu khuyết",
      "Khóa chính", "Khóa ngoại", "Độ mịn (grain)", "Vai trò trong đề án"]
OW = [5, 16, 16, 12, 10, 16, 16, 40, 26, 60]

put_title(ws, "TỪ ĐIỂN DỮ LIỆU — CÁC BẢNG ĐẦU VÀO",
          "Đề án: Ứng dụng mô hình học máy trong hỗ trợ kiểm duyệt tín dụng — "
          "Đồng Quốc Huy (B25CHHT028), Học viện Công nghệ Bưu chính Viễn thông", len(OH))
for i, h in enumerate(OH, start=1):
    ws.cell(row=4, column=i, value=h)
style_header(ws, 4, len(OH))

r = 5
for stt, (name, vn, pk, fk, grain, purpose) in enumerate(TABLES, start=1):
    df = meta[name]
    # Ghi giá trị trực tiếp thay vì công thức: máy chạy không có LibreOffice để
    # tính lại, mà công thức chưa tính sẽ đọc ra rỗng ở pandas và mọi công cụ
    # xem nhanh. Đây là tài liệu tra cứu tĩnh nên giá trị cố định là phù hợp.
    n_null_cols = int((df.isna().mean() > 0).sum())
    vals = [stt, name, vn, len(df), len(df.columns), n_null_cols,
            pk, fk, grain, purpose]
    for i, v in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=i, value=v)
        cell.font = Font(name=FONT, size=10)
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=i in (8, 9, 10),
                                   horizontal="center" if i in (1, 4, 5, 6) else "left")
    ws.cell(row=r, column=4).number_format = "#.##0"
    if stt % 2 == 0:
        for i in range(1, len(OH) + 1):
            ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=C_BAND)
    r += 1

notes = [
    "",
    "QUY ƯỚC CHUNG",
    "• Mốc đúng đắn theo thời điểm: mọi đặc trưng của một hồ sơ chỉ được tính từ dữ liệu phát sinh TRƯỚC applications.apply_time.",
    "• Mọi giá trị tiền tệ tính bằng VND; lãi suất và các tỷ lệ lưu dưới dạng thập phân (0,182 = 18,2%).",
    "• Cột có tên bắt đầu bằng dấu gạch dưới là cột nội bộ của chương trình sinh dữ liệu, KHÔNG dùng làm đầu vào mô hình.",
    "• Dữ liệu khuyết ở nhóm chỉ tiêu CIC phản ánh khách hàng thiếu lịch sử tín dụng (thin-file) — đây là tình huống nghiệp vụ thật, không phải lỗi dữ liệu.",
    "• Biến mục tiêu là labels.default_flag, định nghĩa theo nhóm nợ 3–5 (quá hạn từ 90 ngày) trong cửa sổ quan sát 12 tháng, theo Thông tư 11/2021/TT-NHNN.",
    "",
    "NGUỒN: các file data/raw/*.parquet sinh bởi src/creditreview/data/generator.py; danh mục mã lấy từ src/creditreview/data/reference.py.",
]
for i, t in enumerate(notes):
    c = ws.cell(row=r + i, column=1, value=t)
    c.font = Font(name=FONT, size=10, bold=(t == "QUY ƯỚC CHUNG"), color="1F3864" if t == "QUY ƯỚC CHUNG" else "000000")

for i, w in enumerate(OW, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A5"

# ---- sheet Danh mục mã -----------------------------------------------------
sh = wb.create_sheet("Danh mục mã")
put_title(sh, "DANH MỤC MÃ VÀ GIÁ TRỊ HỢP LỆ",
          "Nguồn: src/creditreview/data/reference.py", 4)
CH = ["Danh mục", "Trường áp dụng", "Mã / Giá trị", "Diễn giải"]
for i, h in enumerate(CH, start=1):
    sh.cell(row=3, column=i, value=h)
style_header(sh, 3, len(CH))

r = 4
for cat, field, items in CODE_LISTS:
    first = True
    for code, expl in items:
        vals = [cat if first else "", field if first else "", code, expl]
        for i, v in enumerate(vals, start=1):
            cell = sh.cell(row=r, column=i, value=v)
            cell.font = Font(name=FONT, size=10, bold=(i == 1 and first))
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(i == 4))
        r += 1
        first = False
    for i in range(1, len(CH) + 1):
        sh.cell(row=r, column=i).border = Border(top=THIN)
    r += 1

for i, w in enumerate([22, 38, 34, 62], start=1):
    sh.column_dimensions[get_column_letter(i)].width = w
sh.freeze_panes = "A4"

wb.save(OUT)
print("Da ghi file.")
